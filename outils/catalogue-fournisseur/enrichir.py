# -*- coding: utf-8 -*-
"""Construit le fichier catalogue que la plateforme attend, images comprises.

Entree  : la tarification du fournisseur + la moisson des fiches publiques.
Sortie  : un XLSX aux colonnes que l'import reconnait, avec une colonne IMAGE_URLS.

La jonction se fait sur la REFERENCE LCM, lue sur la fiche : l'identifiant present
dans l'adresse de la fiche est un numero interne au site, etranger a nos SKU.
"""
import json, re, sys
import openpyxl
from openpyxl import Workbook

# L'import ne retient que les premieres images de la cellule ; autant livrer les
# meilleures. Le site sert plusieurs tailles sous le meme chemin : on demande l'originale.
FILTRE = re.compile(r"/media/cache/(resolve/)?[a-z0-9_]+/")
REMPLACEMENT = "/media/cache/resolve/sylius_shop_product_original/"
MAX_IMAGES = 5


def originale(url):
    return FILTRE.sub(REMPLACEMENT, url, count=1)


def charger_moisson(chemin):
    """ref LCM -> (nom complet, liste d'images) ; la fiche la plus fournie l'emporte."""
    par_ref = {}
    for ligne in open(chemin, encoding="utf-8"):
        f = json.loads(ligne)
        ref = (f.get("ref") or "").strip()
        if not ref:
            continue
        images, vues = [], set()
        for u in f.get("images") or []:
            o = originale(u)
            if o not in vues:
                vues.add(o)
                images.append(o)
        if not images:
            continue
        if ref not in par_ref or len(images) > len(par_ref[ref][1]):
            par_ref[ref] = (f.get("nom"), images)
    return par_ref


def main(tarif, moisson, sortie):
    par_ref = charger_moisson(moisson)
    wb = openpyxl.load_workbook(tarif, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    out = Workbook()
    o = out.active
    o.title = "Catalogue"
    o.append(["SKU", "NOM", "MARQUE", "REFERENCE FOURNISSEUR",
              "OBSERVATIONS", "PU HT", "IMAGE_URLS"])

    lignes = avec_images = sans_fiche = total_images = 0
    for r in ws.iter_rows(min_row=4, values_only=True):
        ref = r[1]
        if ref is None or not str(ref).strip().isdigit():
            continue
        ref = str(ref).strip()
        nom_fichier = (r[3] or "").strip()
        lignes += 1

        fiche = par_ref.get(ref)
        if fiche:
            images = fiche[1][:MAX_IMAGES]
            avec_images += 1
            total_images += len(images)
        else:
            sans_fiche += 1
            images = []

        # Le nom vient TOUJOURS du fichier, jamais de la fiche. Une page du site couvre
        # plusieurs references — les coloris d'un meme fauteuil partagent la meme fiche — et
        # son titre est donc generique. Reprendre ce titre renommerait « NEW BERGEN 1M BEIGE »
        # en « Fauteuil releveur NEW BERGEN », effacant le coloris de produits deja vendus.
        o.append([ref, nom_fichier, r[4], r[5], r[2], r[7], " | ".join(images)])

    out.save(sortie)
    print("lignes          : %d" % lignes)
    print("avec images     : %d" % avec_images)
    print("sans fiche      : %d" % sans_fiche)
    print("images portees  : %d" % total_images)
    print("ecrit           : %s" % sortie)


if __name__ == "__main__":
    main(*sys.argv[1:4])
